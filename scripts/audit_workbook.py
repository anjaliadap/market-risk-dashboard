"""Read-only workbook inventory. Run locally; generated cell inventory stays in data/."""

import argparse
import hashlib
import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from src.settings import settings


def formula_text(cell):
    return str(getattr(cell.value, "text", cell.value))


def chart_ref(obj):
    for key in ("numRef", "strRef"):
        ref = getattr(obj, key, None)
        if ref is not None:
            return ref.f
    return None


def chart_text(obj):
    if obj is None:
        return ""
    return " ".join(n.text for n in obj.to_tree().iter() if n.tag.endswith("}t") and n.text)


def audit(path):
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    report = {"sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest(), "sheets": []}
    report["external_links"] = len(formulas._external_links)
    for sheet in formulas:
        cached = values[sheet.title]
        cells, counts = [], Counter()
        for row in sheet:
            for cell in row:
                if cell.value is None:
                    continue
                value = cached[cell.coordinate].value
                is_formula = cell.data_type == "f"
                text = formula_text(cell) if is_formula else None
                kind = "formula" if is_formula else "hardcoded"
                if text and any(t in text.upper() for t in ("BDH(", "BDP(", "BQL(", "BDS(")):
                    kind = "bloomberg"
                counts[kind] += 1
                if is_formula and value is None:
                    counts["missing_cache"] += 1
                if cached[cell.coordinate].data_type == "e" or (
                    isinstance(value, str) and value.startswith(("#N/A", "#REF!", "#VALUE!"))
                ):
                    counts["invalid"] += 1
                if isinstance(value, (date, datetime)):
                    counts["dated_cells"] += 1
                color = cell.font.color
                cells.append(
                    {
                        "cell": cell.coordinate,
                        "formula": text,
                        "value": value,
                        "kind": kind,
                        "format": cell.number_format,
                        "font_color": str(color.rgb) if color and color.type == "rgb" else None,
                    }
                )
        charts = []
        for index, chart in enumerate(sheet._charts):
            series = []
            for subchart in chart._charts:
                for item in subchart.series:
                    series.append(
                        {
                            "x": chart_ref(getattr(item, "cat", None))
                            or chart_ref(getattr(item, "xVal", None)),
                            "y": chart_ref(getattr(item, "val", None))
                            or chart_ref(getattr(item, "yVal", None)),
                            "name_ref": chart_ref(getattr(item, "tx", None)),
                            "name": getattr(getattr(item, "tx", None), "v", None),
                            "type": subchart.__class__.__name__,
                            "axis": getattr(getattr(subchart, "y_axis", None), "axId", None),
                        }
                    )
            charts.append({"index": index, "title": chart_text(chart.title), "series": series})
        report["sheets"].append(
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "counts": dict(counts),
                "charts": charts,
                "cells": cells,
            }
        )
    formulas.close()
    values.close()
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument(
        "--output", type=Path, default=Path(settings.excel_artifact_dir) / "workbook_audit.json"
    )
    args = parser.parse_args()
    result = audit(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, default=str, indent=2), encoding="utf-8")
    print(f"Audited {len(result['sheets'])} sheets; local inventory: {args.output}")
