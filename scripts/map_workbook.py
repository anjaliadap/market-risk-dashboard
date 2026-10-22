"""Build a reviewable layout map from the local audit, never copying cached financial data."""

import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl.utils.cell import get_column_letter, range_boundaries

from src.settings import settings

MANUAL = {
    "Market Data": "Authorized Bloomberg refresh/save; snapshot TODAY() dates are calculation dates, not release dates. Sector breadth is constituent breadth, not index distance to DMA. Futures prices must not be labeled spot.",
    "Bear Market": "D8:D21 current, E8:E21 historical averages (linked to row 22, preserve them), G8:G21 explicit direction rules; F8:F21 saved flags. D23=14 denominator. Historical peaks and IPO inputs are sponsor assumptions. Never infer red flags from font colors.",
    "S&P vs Bonds": "T:AA Bloomberg inputs; J:R cached Excel outputs. P/S yield less Baa yield and equity P/E less Baa P/E are sponsor definitions, not an equity risk premium.",
    "Valuation": "C:E are fixed historical benchmarks; F current Bloomberg values; H current vs 2000. Preserve benchmark dates and values.",
    "S&P 500 Fundamentals": "B:R Bloomberg/BQL sources; W:AC sponsor calculations. L23 and Q23 have no formula cache (BQL spill anchors); downstream dated numeric results are independent saved values. Large BQL tables include identity/header rows, not observations.",
    "Margin Debt": "W:X GDP, Z:AA FINRA margin debt in millions, AC:AD Wilshire market cap in millions. FINRA data must be manually updated; market cap hardcoded through early 2026 then Bloomberg. K:U retains sponsor unit alignment/calculations. B:G peak/trough benchmarks preserved.",
    "S&P Relative": "P:Q index, S:T gold, V:W WTI. WTI hardcoded through 1999, individually dated formulas thereafter; authorized operator must append next month and extend B:N calculations. Sheet explicitly says 7-year MA/difference need updating. Do not silently repair rolling windows.",
    "Gold to Oil": "B:F Bloomberg inputs, H:L cached ratios/long-term average; N23:Q29 implied prices and upside/downside use fixed row 463. New observations require extending formulas AND updating these fixed endpoints in Excel.",
    "Bull vs Bear": "B:F and N:Q Bloomberg; H:L Excel; S:T fixed decile lookup; V:AA saved conditional returns. Forward returns contain BQL capacity-error text. Unavailable values stay missing, never zero. Historical conditional returns are descriptive and overlap/publication controls are not established.",
    "AAII Bull": "Auxiliary substantive worksheet, not a separate sponsor section. Additional Bloomberg returns and hardcoded decile lookup; mapped as an optional Bull vs Bear chart. Do not substitute it for the requested bull-minus-bear study.",
    "VIX & Options": "Bloomberg series plus cached current-reference lines. Chart endpoints stop at row 1078 although sheet extends to row 1399; reproduce configured chart range, report clipping, and require mapping review to extend.",
    "BTC": "D4 Bitcoin position, D8 debt/preferred sum, D15 cost basis manually maintained from Strategy Analytics/Ledger. D5 Bloomberg BTC price; D6:D16 cached sponsor outputs; USD millions vs USD/BTC preserved.",
    "Banks": "Bloomberg quarterly net charge-offs; population/units are from sponsor workbook, do not replace with another bank universe.",
    "Housing": "Independent date columns for different monthly/quarterly frequencies. Preserve secondary axes; never align affordability with the longer mortgage series by row position.",
    "Energy": "Bloomberg source pairs B:C, E:F, H:I, N:O, Q:R, T:U, W:X, Z:AA. K:L and AC:AD are calculated efficiencies. WTI/Henry Hub use futures tickers CL1/NG1; gas Bcfd and Bcf axes differ. Chart endpoints are shorter than source columns.",
    "Recession": "B/D/F/H checklist has manual text states and Bloomberg numbers; G saved flags; D55=17 denominator. F55 SUM(G8:G55) includes total row but G55 is blank. H notes govern thresholds. U Sahm warning line is hardcoded 0.5. Manual statuses are not proven current by a new save.",
    "Consumer": "Bloomberg series and source-specific units; debt-service rate and savings rate have different denominators.",
    "Demand": "Horizontal quarterly history/forward estimates; row 2 labels distinguish (Rep)/(Fwd), row 3 dates. Preserve red vertical latest-reported-quarter marker. Revenue/FCF/capex Excel calculations remain authoritative; do not describe all revenue as AI-specific.",
    "Financing": "Sponsor Financing section is additional to existing Forecasting. Horizontal hyperscaler ratios and combined quarterly cash flows, plus credit/relative-performance histories. Forward estimates must be distinguished from reported quarters.",
}
TABLES = {
    "Market Data": [("market_snapshot", "B3:M54", "Z4")],
    "Bear Market": [("bear_checklist", "B7:G26", None)],
    "Valuation": [("valuation_benchmarks", "B2:H12", "F3")],
    "Margin Debt": [("margin_benchmarks", "B4:G16", None)],
    "Gold to Oil": [("implied_prices", "N23:Q29", "H463")],
    "Bull vs Bear": [("sentiment_deciles", "V3:AA16", None)],
    "BTC": [("btc_coverage", "B3:E16", None)],
    "Recession": [("recession_checklist", "B6:H56", None)],
}
TABLE_COLUMNS = {
    "market_snapshot": [
        ("B", "Market"),
        ("D", "Current"),
        ("E", "YTD"),
        ("F", "1M"),
        ("G", "3M"),
        ("H", "6M"),
        ("I", "12M"),
        ("J", "50-DMA"),
        ("K", "200-DMA"),
        ("L", "52W Low"),
        ("M", "52W High"),
    ],
    "bear_checklist": [("B", "Indicator"), ("D", "Current"), ("E", "Historical average")],
    "valuation_benchmarks": [
        ("B", "Valuation metric"),
        ("C", "2000 peak"),
        ("D", "2002 trough"),
        ("E", "2009 trough"),
        ("F", "Current"),
        ("H", "Current vs 2000 peak"),
    ],
    "margin_benchmarks": [
        ("B", "Period"),
        ("D", "Peak"),
        ("E", "Trough"),
        ("F", "Prior peak"),
        ("G", "Prior trough"),
    ],
    "implied_prices": [("N", "Metric"), ("Q", "Saved value")],
    "sentiment_deciles": [
        ("V", "Investors"),
        ("W", "Bull less bear decile"),
        ("X", "1M"),
        ("Y", "3M"),
        ("Z", "6M"),
        ("AA", "12M"),
    ],
    "btc_coverage": [("B", "Metric"), ("D", "Saved value")],
    "recession_checklist": [
        ("B", "Metric"),
        ("D", "Typical recession reading"),
        ("F", "Current status"),
        ("H", "Sponsor notes"),
    ],
}


def split_ref(ref):
    sheet, cells = ref.rsplit("!", 1)
    return sheet.strip("'").replace("''", "'"), cells.replace("$", "")


def build(audit):
    mapping = {
        "version": 1,
        "policy": "Saved caches only; no Bloomberg refresh attestation",
        "sections": {},
    }
    details = []
    for sheet in audit["sheets"]:
        name = sheet["name"]
        if name not in MANUAL:
            continue
        section = "Bull vs Bear" if name == "AAII Bull" else name
        config = mapping["sections"].setdefault(
            section, {"sheets": [], "notes": [], "charts": [], "tables": []}
        )
        config["sheets"].append(name)
        config["notes"].append(MANUAL[name])
        cells = {c["cell"]: c for c in sheet["cells"]}
        for ident, area, observed in TABLES.get(name, []):
            a, first, b, _ = range_boundaries(area)
            headers = {
                f"{get_column_letter(c)}{first}": cells[f"{get_column_letter(c)}{first}"]["value"]
                for c in range(a, b + 1)
                if isinstance(cells.get(f"{get_column_letter(c)}{first}", {}).get("value"), str)
            }
            config["tables"].append(
                {
                    "id": ident,
                    "title": ident.replace("_", " ").title(),
                    "sheet": name,
                    "range": area,
                    "date_cell": observed,
                    "expected_headers": headers,
                    "display_columns": TABLE_COLUMNS[ident],
                    "signal_column": "F"
                    if name == "Bear Market"
                    else "G"
                    if name == "Recession"
                    else None,
                }
            )
        for chart in sheet["charts"]:
            chart_id = name.lower().replace(" ", "_").replace("&", "and") + f"_{chart['index']}"
            out = {
                "id": chart_id,
                "title": chart["title"] or f"{name} chart {chart['index'] + 1}",
                "series": [],
            }
            axis_ids = []
            for idx, series in enumerate(chart["series"]):
                if not series["y"]:
                    continue
                ys, yr = split_ref(series["y"])
                c1, r1, _c2, r2 = range_boundaries(yr)
                xr = series["x"]
                inferred = False
                if not xr:
                    # Excel combo charts omit categories on secondary groups. Their own
                    # adjacent date column is essential when frequency/length differs.
                    for c in range(c1 - 1, 0, -1):
                        val = cells.get(f"{get_column_letter(c)}{r1}", {}).get("value")
                        try:
                            datetime.fromisoformat(str(val))
                        except (ValueError, TypeError):
                            continue
                        xr = f"'{ys}'!{get_column_letter(c)}{r1}:{get_column_letter(c)}{r2}"
                        inferred = True
                        break
                if not xr:
                    raise ValueError(f"Unresolved date range: {name} {yr}")
                xs, xcells = split_ref(xr)
                label = series["name"]
                if series["name_ref"]:
                    _, nc = split_ref(series["name_ref"])
                    a, b, _, _ = range_boundaries(nc)
                    label = cells.get(f"{get_column_letter(a)}{b}", {}).get("value")
                label = (
                    label
                    or cells.get(f"{get_column_letter(c1)}{r1 - 1}", {}).get("value")
                    or f"Series {idx + 1}"
                )
                if series["axis"] not in axis_ids:
                    axis_ids.append(series["axis"])
                if not isinstance(label, str):
                    label = "Combined" if name == "Financing" else f"Series {idx + 1}"
                formats = Counter(
                    cells.get(f"{get_column_letter(col)}{row}", {}).get("format")
                    for row in range(r1, r2 + 1)
                    for col in range(c1, _c2 + 1)
                )
                formats.pop(None, None)
                fmt = formats.most_common(1)[0][0] if formats else "General"
                populated = [
                    cells.get(f"{get_column_letter(col)}{row}", {}).get("value")
                    for row in range(r1, r2 + 1)
                    for col in range(c1, _c2 + 1)
                ]
                marker = (
                    r1 == r2
                    and "Bar" in series["type"]
                    and sum(isinstance(v, (int, float)) and v != 0 for v in populated) <= 1
                )
                if marker:
                    label = "Latest reported quarter"
                out["series"].append(
                    {
                        "id": chart_id + f"_{idx}",
                        "name": str(label).replace("\n", " "),
                        "role": "reported_marker" if marker else "metric",
                        "sheet": ys,
                        "values": yr,
                        "date_sheet": xs,
                        "dates": xcells,
                        "date_mapping_inferred": inferred,
                        "kind": series["type"],
                        "secondary": axis_ids.index(series["axis"]) > 0,
                        "unit": fmt,
                        "value_type": "number",
                        "date_type": "excel_date",
                        "expected_number_format": fmt,
                        "allowed_number_formats": list(formats),
                        "allow_forward": name in ["Demand", "Financing"] and r1 == r2,
                        "period_labels_row": 2 if r1 == r2 else None,
                    }
                )
            config["charts"].append(out)
        details.append((sheet, config))
    return mapping, details


def write_docs(audit, mapping, details):
    lines = [
        "# Excel integration audit and mapping",
        "",
        "Audit completed before provider implementation. The supplied Word document is the authority for chart intent and layout; workbook chart references are the authority for saved series. No source workbook is modified. Full cell/formula/cache/font inventory is local-only in `EXCEL_ARTIFACT_DIR/workbook_audit.json` (defaults to local AppData, outside the synchronized project); no cached financial values or document images are committed.",
        "",
        f"Workbook SHA-256 at audit: `{audit['sha256']}`. {len(audit['sheets'])} worksheets, 19 substantive sheets, no external workbook links. Divider/cache sheets: Market>, Macro>, AI>, PB_CACHE_JS. Bloomberg XLL references are add-in calls, not external-workbook links.",
        "",
        "## Findings and refresh contract",
        "",
        "- Almost all formulas have saved caches. Fundamentals L23/Q23 have no cache. Bull vs Bear has 2,025 error-text cells in its forward-return area, despite no native Excel error-typed cells. Errors encoded as text must also be rejected.",
        "- Cached values establish what Excel last saved, not when Bloomberg was refreshed. TODAY(), file name, modification time and successful import are never refresh attestations. Several dates extend into forward estimate periods; they are not latest observed dates.",
        "- Readiness requires authorized Excel/Bloomberg refresh, formula extension where noted, recalculation and save. Python never evaluates XLL/BQL/BDH/BDP or Excel formulas. Missing formula caches remain unavailable.",
        "- Chart ranges are deliberately explicit. Source spills and worksheet used ranges often extend past chart endpoints; auto-extending these could change historical benchmarks or forecasts. Central mapping review is required after layout/endpoint changes.",
        "- Tables retain raw sponsor values and number formats. Percent fractions are displayed using their Excel formats, not guessed scaling. Undated snapshots remain explicitly undated. Native errors, text errors, missing caches and manual inputs are reported.",
        "- Recession and bear-market saved checklist flags may lag manual states. Saved scores are displayed as sponsor outputs, not recalibrated predictions. Do not infer financial signals from red font (which also denotes manual inputs).",
        "",
        "## Section coverage",
        "",
        "Existing Forecasting remains an assumption tool. Add Financing (present in sponsor files) and Supply (no workbook sheet) without removing existing pages. Overview summarizes section coverage. FRED remains independent fallback; Treasury, SEC and EIA are not currently implemented connectors.",
        "",
        "`config/excel_mapping.json` is the executable, versioned contract. Each chart series has a named metric ID, sheet, values/date ranges, units/number format, type and axis. Missing category references in combo charts are resolved to their own adjacent date columns, not borrowed from a differently sampled series. These inferences are marked for review.",
        "",
    ]
    for sheet, _ in details:
        name = sheet["name"]
        lines += [
            f"## {name}",
            "",
            f"Used range: A1:{get_column_letter(sheet['columns'])}{sheet['rows']}. Inventory: `{json.dumps(sheet['counts'])}`.",
            "",
            MANUAL[name],
            "",
            "### Charts and source ranges",
            "",
            "| Chart | Metric | Saved values | Dates | Type / axis |",
            "|---|---|---|---|---|",
        ]
        section = mapping["sections"]["Bull vs Bear" if name == "AAII Bull" else name]
        for chart in section["charts"]:
            for series in chart["series"]:
                if series["sheet"] != name:
                    continue
                lines.append(
                    f"| {chart['title']} | {series['name']} | {series['values']} | {series['dates']}{' (inferred own dates)' if series['date_mapping_inferred'] else ''} | {series['kind']} / {'right' if series['secondary'] else 'left'} |"
                )
        lines += [
            "",
            "### Formula and manual-extension inventory by column",
            "",
            "Last formula vs last nonempty row exposes gaps and fixed endpoints; this is a review aid, not an instruction to fill cells.",
            "",
            "| Column | Formula count | First / last formula | Last nonempty | Example formula |",
            "|---|---:|---|---:|---|",
        ]
        columns = {}
        for cell in sheet["cells"]:
            col = "".join(filter(str.isalpha, cell["cell"]))
            group = columns.setdefault(col, [])
            group.append(cell)
        for col, group in columns.items():
            fs = [c for c in group if c["formula"]]
            if fs:
                lines.append(
                    f"| {col} | {len(fs)} | {fs[0]['cell']} / {fs[-1]['cell']} | {group[-1]['cell']} | `{fs[0]['formula'][:180].replace('|', '/').replace(chr(10), ' ')}` |"
                )
        reds = [
            c["cell"]
            for c in sheet["cells"]
            if c["kind"] == "hardcoded" and c["font_color"] in ["FFFF0000", "00FF0000"]
        ]
        lines += [
            "",
            f"Red-font hardcoded cells: {len(reds)}; examples `{', '.join(reds[:35])}`. All constants and full formulas are retained in the local inventory for operator review.",
            "",
        ]
    lines += [
        "## Supply and independent-source plan",
        "",
        "No workbook worksheet exists for Supply. Show unavailable coverage explicitly. Candidate capacity/vacancy: company filings and operator disclosures (keep live vs construction, geography, MW and measurement dates separate). GPU rental: provider-published rate cards/APIs with model, region, VRAM, availability, hourly price, currency and on-demand/spot/contract terms; do not combine these contracts into one price. Bottlenecks: dated utility/interconnection and manufacturer disclosures. No verified connector or permission to scrape is implied. SEC facts can supplement filed capex/CFO; EIA can supplement energy; FRED can supply productivity. Each requires a separate verified connector and attribution before use. No synthetic observations.",
        "",
        "## Visual acceptance and limitations",
        "",
        "Word contains 31 chart/table illustrations, reviewed locally. Plum primary series, teal comparison, secondary axes, bar/line combinations, horizontal company grids, and last-reported-quarter red markers are mapped into interactive Plotly charts. Pixel-identical Excel rendering is not claimed. Original chart endpoints and cached calculation semantics are preserved. Some chart titles are blank in Excel; metadata supplies descriptive titles. Conditional formatting rules/checklist operator judgments are not recalculated. Formula extensions, Bloomberg freshness and independent Supply data remain operator/external dependencies.",
    ]
    Path("docs/excel_integration_map.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    audit = json.loads(
        (Path(settings.excel_artifact_dir) / "workbook_audit.json").read_text(encoding="utf-8")
    )
    mapping, details = build(audit)
    # Blank chart titles are presentation metadata, not source-data changes.
    names = {
        "S&P vs Bonds": ["P/S yield less Baa bond yield", "Equity P/E less Baa bond P/E"],
        "Housing": {2: "Existing home median price", 5: "Homeowners' equity"},
        "Energy": {0: "US oil production"},
        "Consumer": [
            "Household debt service ratio",
            "Consumer credit outstanding",
            "Personal savings rate",
        ],
    }
    for section, titles in names.items():
        for i, chart in enumerate(mapping["sections"][section]["charts"]):
            if isinstance(titles, list) or i in titles:
                chart["title"] = titles[i]
    Path("config/excel_mapping.json").write_text(
        json.dumps(mapping, indent=2) + "\n", encoding="utf-8"
    )
    write_docs(audit, mapping, details)
    print(
        f"Mapped {sum(len(s['charts']) for s in mapping['sections'].values())} charts across {len(mapping['sections'])} sections; audit written first."
    )
