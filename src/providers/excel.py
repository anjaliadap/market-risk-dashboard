"""Read saved Excel values. Never evaluate formulas or save the source workbook."""

import hashlib
import json
import math
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import from_excel

MAPPING_PATH = Path(__file__).resolve().parents[2] / "config/excel_mapping.json"
ERROR_PREFIXES = (
    "#N/A",
    "#REF!",
    "#VALUE!",
    "#DIV/0!",
    "#NAME?",
    "#NUM!",
    "#SPILL!",
    "#CALC!",
    "#GETTING_DATA",
    "#N/A Request",
    "#BLOCKED!",
)


def load_mapping():
    return json.loads(MAPPING_PATH.read_text(encoding="utf-8"))


def is_error(value):
    return isinstance(value, str) and value.lstrip().startswith(ERROR_PREFIXES)


def normalize_date(value, epoch):
    if isinstance(value, bool) or value is None or is_error(value):
        raise ValueError("Missing or invalid observation date")
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or not 1 <= value <= 100000:
            raise ValueError("Invalid Excel serial date")
        value = from_excel(value, epoch)
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return datetime.fromisoformat(value).date().isoformat()
    raise ValueError("Unsupported observation date")


def numeric(value, spec):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError("Missing, nonnumeric or nonfinite saved value")
    if "min" in spec and value < spec["min"]:
        raise ValueError("Value below configured unit bound")
    if "max" in spec and value > spec["max"]:
        raise ValueError("Value above configured unit bound")
    return float(value)


def range_cells(sheet, area):
    c1, r1, c2, r2 = range_boundaries(area)
    if c1 != c2 and r1 != r2:
        raise ValueError("Time-series ranges must be one row or one column")
    return [
        cell
        for row in sheet.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2)
        for cell in row
    ]


class ExcelWorkbookProvider:
    def __init__(self, path, mapping=None):
        self.path = Path(path).expanduser().resolve()
        self.mapping = mapping if mapping is not None else load_mapping()

    def read(self):
        before = self.path.stat()
        blob = self.path.read_bytes()
        after = self.path.stat()
        if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
            raise ValueError("Workbook changed during read; wait for Excel save to finish")
        return self.read_bytes(blob, after.st_mtime)

    def read_bytes(self, blob, modified_at=None):
        # Both views come from the exact same immutable byte buffer.
        formulas = load_workbook(BytesIO(blob), data_only=False, keep_links=True)
        cached = load_workbook(BytesIO(blob), data_only=True, keep_links=True)
        try:
            return self._extract(formulas, cached, blob, modified_at)
        finally:
            formulas.close()
            cached.close()

    def _extract(self, formulas, cached, blob, modified_at):
        issues = []
        diagnostics = {"missing_formula_caches": 0, "failed_bloomberg_values": 0, "error_values": 0}

        def cell_value(sheet, coordinate):
            cell = cached[sheet][coordinate]
            formula = formulas[sheet][coordinate].data_type == "f"
            value = cell.value
            problem = None
            if formula and value is None:
                problem = "Formula has no saved cache"
            elif cell.data_type == "e" or is_error(value):
                problem = "Excel/Bloomberg error in saved value"
            if problem:
                issues.append({"sheet": sheet, "cell": coordinate, "reason": problem})
                value = None
            return value, formula

        required_sheets = {
            s for section in self.mapping["sections"].values() for s in section["sheets"]
        }
        missing = required_sheets - set(cached.sheetnames)
        if missing:
            raise ValueError(f"Missing mapped worksheets: {', '.join(sorted(missing))}")
        for name in required_sheets:
            for row in formulas[name]:
                for cell in row:
                    val = cached[name][cell.coordinate]
                    if cell.data_type == "f" and val.value is None:
                        diagnostics["missing_formula_caches"] += 1
                    if val.data_type == "e" or is_error(val.value):
                        diagnostics["error_values"] += 1
                        if cell.data_type == "f" or "BQL" in str(val.value):
                            diagnostics["failed_bloomberg_values"] += 1
        sections = {}
        latest_dates = []
        metric_count = 0
        for section_name, config in self.mapping["sections"].items():
            section = {"notes": config["notes"], "charts": [], "tables": []}
            for table in config["tables"]:
                for coord, expected in table.get("expected_headers", {}).items():
                    if cached[table["sheet"]][coord].value != expected:
                        raise ValueError(f"Layout changed: {table['sheet']}!{coord}")
                c1, r1, c2, r2 = range_boundaries(table["range"])
                rows = []
                for row in cached[table["sheet"]].iter_rows(
                    min_row=r1, max_row=r2, min_col=c1, max_col=c2
                ):
                    output = []
                    for cell in row:
                        value, formula = cell_value(table["sheet"], cell.coordinate)
                        if isinstance(value, float) and not math.isfinite(value):
                            issues.append(
                                {
                                    "sheet": table["sheet"],
                                    "cell": cell.coordinate,
                                    "reason": "Nonfinite table value",
                                }
                            )
                            value = None
                        if isinstance(value, (date, datetime)):
                            value = value.isoformat()[:10]
                        output.append(
                            {
                                "cell": cell.coordinate,
                                "value": value,
                                "format": cell.number_format,
                                "formula": formula,
                                "font_color": str(
                                    formulas[table["sheet"]][cell.coordinate].font.color.rgb
                                )
                                if formulas[table["sheet"]][cell.coordinate].font.color
                                and formulas[table["sheet"]][cell.coordinate].font.color.type
                                == "rgb"
                                else None,
                            }
                        )
                    rows.append(output)
                observed = None
                if table.get("date_cell"):
                    value, _ = cell_value(table["sheet"], table["date_cell"])
                    if value is not None:
                        observed = normalize_date(value, cached.epoch)
                        latest_dates.append(observed)
                section["tables"].append({**table, "rows": rows, "as_of": observed})
                metric_count += sum(
                    isinstance(c["value"], (int, float)) for row in rows for c in row
                )
            for chart in config["charts"]:
                chart_out = {"id": chart["id"], "title": chart["title"], "series": []}
                for spec in chart["series"]:
                    if not spec.get("unit") or spec.get("value_type") != "number":
                        raise ValueError(f"Metric unit/type contract missing: {spec['id']}")
                    dates = range_cells(cached[spec["date_sheet"]], spec["dates"])
                    values = range_cells(cached[spec["sheet"]], spec["values"])
                    if len(dates) != len(values):
                        raise ValueError(f"Date/value range length mismatch: {spec['id']}")
                    points, seen = [], set()
                    reported_dates = []
                    for dc, vc in zip(dates, values, strict=True):
                        d, _ = cell_value(spec["date_sheet"], dc.coordinate)
                        v, _ = cell_value(spec["sheet"], vc.coordinate)
                        if d is None and v is None:
                            continue  # Explicit chart ranges may contain trailing empty rows.
                        expected_format = spec.get("expected_number_format")
                        allowed_formats = spec.get("allowed_number_formats", [expected_format])
                        if (
                            v is not None
                            and expected_format
                            and spec.get("role") != "reported_marker"
                            and vc.number_format not in allowed_formats
                        ):
                            raise ValueError(
                                f"Unit/number format changed: {spec['sheet']}!{vc.coordinate}"
                            )
                        try:
                            normalized = normalize_date(d, cached.epoch)
                        except ValueError as exc:
                            raise ValueError(
                                f"Invalid date {spec['date_sheet']}!{dc.coordinate}"
                            ) from exc
                        if normalized in seen:
                            raise ValueError(f"Duplicate date in {spec['id']}: {normalized}")
                        seen.add(normalized)
                        try:
                            v = numeric(v, spec)
                        except ValueError as exc:
                            if spec.get("role") != "reported_marker":
                                issues.append(
                                    {
                                        "sheet": spec["sheet"],
                                        "cell": vc.coordinate,
                                        "reason": str(exc),
                                    }
                                )
                            v = None
                        forward = False
                        if spec.get("period_labels_row"):
                            label = (
                                cached[spec["date_sheet"]]
                                .cell(spec["period_labels_row"], dc.column)
                                .value
                            )
                            if not isinstance(label, str) or not any(
                                t in label for t in ("(Rep)", "(Fwd)")
                            ):
                                raise ValueError(
                                    f"Missing reported/forward label for {spec['id']} {dc.coordinate}"
                                )
                            forward = "(Fwd)" in label
                            if not forward:
                                reported_dates.append(normalized)
                        if v is not None and not forward:
                            latest_dates.append(normalized)
                        source_cell = formulas[spec["sheet"]][vc.coordinate]
                        formula = (
                            str(getattr(source_cell.value, "text", source_cell.value))
                            if source_cell.data_type == "f"
                            else None
                        )
                        points.append(
                            {
                                "date": normalized,
                                "value": v,
                                "cell": vc.coordinate,
                                "forward": forward,
                                "formula": formula,
                            }
                        )
                    points.sort(key=lambda p: p["date"])
                    chart_out["series"].append(
                        {
                            **spec,
                            "points": points,
                            "latest_reported": max(reported_dates, default=None),
                        }
                    )
                    if spec.get("role") != "reported_marker":
                        metric_count += any(p["value"] is not None for p in points)
                section["charts"].append(chart_out)
            sections[section_name] = section
        if not metric_count:
            raise ValueError("No valid mapped numeric metrics in workbook")
        return {
            "sha256": hashlib.sha256(blob).hexdigest(),
            "mapping_sha256": hashlib.sha256(
                json.dumps(self.mapping, sort_keys=True).encode()
            ).hexdigest(),
            "path": str(self.path),
            "modified_at": modified_at,
            "sections": sections,
            "metric_count": metric_count,
            "latest_observation": max(latest_dates, default=None),
            "diagnostics": diagnostics,
            "issues": issues,
            "refresh_attestation": "Unknown: saved workbook caches do not attest Bloomberg refresh",
        }
